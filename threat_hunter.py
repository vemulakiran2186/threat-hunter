import re
import requests
from bs4 import BeautifulSoup
import urllib3
import pandas as pd
import io
import json
from datetime import datetime, timedelta
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from urllib.parse import urlparse, urljoin

# Disable SSL verification warnings to handle self-signed certs gracefully
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

# Regex for matching CVE IDs (e.g. CVE-2023-12345)
CVE_REGEX = re.compile(r'\bCVE-\d{4}-\d{4,7}\b', re.IGNORECASE)

def fetch_url(url, timeout=15):
    """
    Fetches the HTML content of a URL with common browser headers to minimize blocks.
    """
    headers = {
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
        'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,image/apng,*/*;q=0.8',
        'Accept-Language': 'en-US,en;q=0.9',
        'Cache-Control': 'max-age=0',
        'Upgrade-Insecure-Requests': '1'
    }
    
    # Simple URL sanitization
    url = url.strip()
    if not url.startswith(('http://', 'https://')):
        url = 'https://' + url

    try:
        response = requests.get(url, headers=headers, timeout=timeout, verify=False)
        response.raise_for_status()
        return response.text
    except requests.exceptions.RequestException as e:
        raise Exception(f"Failed to connect: {str(e)}")

def is_article_url(url, base_url, link_text=""):
    """
    Heuristic check to determine if a URL represents an individual article/advisory page
    rather than a category page, contact page, or main site index.
    """
    parsed_base = urlparse(base_url)
    parsed_url = urlparse(url)
    
    # Must be same domain (or subdomain)
    base_domain = parsed_base.netloc.replace('www.', '')
    url_domain = parsed_url.netloc.replace('www.', '')
    
    if url_domain and base_domain not in url_domain:
        return False
        
    path = parsed_url.path.lower()
    
    # Skip common static pages and Blogger static directories
    skip_keywords = [
        '/contact', '/about', '/search', '/category', '/tag', '/author', 
        '/privacy', '/terms', '/help', '/rss', '/p/'
    ]
    if any(kw in path for kw in skip_keywords):
        return False
        
    # Check if URL looks like an article
    # 1. Contains a year pattern like /2026/06/
    if re.search(r'/(202[4-9])([-/\d]+)?/', path):
        return True
        
    # 2. Ends with .html and has a slug
    if path.endswith('.html') and '-' in path:
        return True
        
    # 3. Path contains standard article words and has a slug
    article_paths = ['/news/', '/alerts/', '/blog/', '/article/', '/posts/', '/advisories/', '/serious-security/']
    if any(ap in path for ap in article_paths) and len(path.split('/')) > 2:
        if '-' in path or '_' in path:
            return True
            
    # 4. Long link text (likely a headline) with a slug
    if len(link_text.strip().split()) >= 4 and ('-' in path or '_' in path):
        return True
        
    return False

def extract_article_links(base_url, html_content):
    """
    Extracts all candidate article/advisory URLs from a homepage or listing page.
    """
    soup = BeautifulSoup(html_content, 'html.parser')
    links = []
    for a in soup.find_all('a', href=True):
        href = a['href']
        text = a.get_text()
        full_url = urljoin(base_url, href)
        if is_article_url(full_url, base_url, text):
            links.append(full_url)
    return list(dict.fromkeys(links)) # deduplicate

def clean_text(text):
    """
    Cleans up whitespace and formatting in parsed text.
    """
    if not text:
        return ""
    # Replace multiple spaces/newlines with a single space
    text = re.sub(r'\s+', ' ', text)
    return text.strip()

def extract_meta_description(soup):
    """
    Attempts to extract meta descriptions from the HTML.
    """
    meta_desc = soup.find('meta', attrs={'name': 'description'}) or \
                soup.find('meta', attrs={'property': 'og:description'}) or \
                soup.find('meta', attrs={'name': 'twitter:description'})
    
    if meta_desc and meta_desc.get('content'):
        return clean_text(meta_desc.get('content'))
    return ""

def find_cve_context(soup, cve_id):
    """
    Looks for the paragraph or text block containing the specific CVE ID to provide context.
    """
    cve_lower = cve_id.lower()
    # Find all text blocks
    blocks = soup.find_all(['p', 'li', 'td', 'div'])
    for block in blocks:
        # Avoid huge blocks like body or main
        if block.name == 'div' and len(block.get_text()) > 1000:
            continue
            
        block_text = block.get_text()
        if cve_lower in block_text.lower():
            cleaned = clean_text(block_text)
            if 30 <= len(cleaned) <= 600:
                return cleaned
            elif len(cleaned) > 600:
                # Truncate nicely if it's too long
                return cleaned[:597] + "..."
    return ""

def parse_natural_date(text):
    """
    Parses strings like 'June 5, 2026', '05 Jun 2026', '2026-06-05', etc.
    """
    text = text.strip()
    m = re.search(r'\b(\d{4})-(\d{2})-(\d{2})\b', text)
    if m:
        try:
            return datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            pass
            
    months = ['january', 'february', 'march', 'april', 'may', 'june', 
              'july', 'august', 'september', 'october', 'november', 'december',
              'jan', 'feb', 'mar', 'apr', 'may', 'jun', 'jul', 'aug', 'sep', 'oct', 'nov', 'dec']
              
    m = re.search(r'\b([a-zA-Z]{3,9})\s+(\d{1,2}),?\s+(\d{4})\b', text)
    if m:
        month_name = m.group(1).lower()
        if month_name in months:
            idx = months.index(month_name)
            if idx >= 12:
                idx = idx - 12
            month_num = idx + 1
            try:
                return datetime(int(m.group(3)), month_num, int(m.group(2)))
            except ValueError:
                pass
                
    m = re.search(r'\b(\d{1,2})\s+([a-zA-Z]{3,9})\s+(\d{4})\b', text)
    if m:
        month_name = m.group(2).lower()
        if month_name in months:
            idx = months.index(month_name)
            if idx >= 12:
                idx = idx - 12
            month_num = idx + 1
            try:
                return datetime(int(m.group(3)), month_num, int(m.group(1)))
            except ValueError:
                pass
                
    return None

def extract_article_date(soup, url):
    """
    Attempts to extract the publication or modification date from the article page.
    Returns a datetime object, or None if no date can be parsed.
    """
    # 1. Check JSON-LD
    for script in soup.find_all('script', type='application/ld+json'):
        if not script.string:
            continue
        try:
            data = json.loads(script.string)
            dates = []
            def find_dates(obj):
                if isinstance(obj, dict):
                    for k, v in obj.items():
                        if k in ['datePublished', 'dateModified', 'uploadDate']:
                            if isinstance(v, str):
                                dates.append(v)
                        else:
                            find_dates(v)
                elif isinstance(obj, list):
                    for item in obj:
                        find_dates(item)
            find_dates(data)
            for date_str in dates:
                m = re.match(r'^(\d{4}-\d{2}-\d{2})', date_str)
                if m:
                    return datetime.strptime(m.group(1), "%Y-%m-%d")
        except Exception:
            pass

    # 2. Check itemprop attributes
    for tag in soup.find_all(attrs={'itemprop': re.compile(r'date', re.IGNORECASE)}):
        content = tag.get('content') or tag.get_text()
        if content:
            m = re.match(r'^(\d{4}-\d{2}-\d{2})', content.strip())
            if m:
                try:
                    return datetime.strptime(m.group(1), "%Y-%m-%d")
                except ValueError:
                    pass
            parsed_dt = parse_natural_date(content)
            if parsed_dt:
                return parsed_dt

    # 3. Check meta tags
    meta_tags = [
        ('property', 'article:published_time'),
        ('property', 'article:modified_time'),
        ('name', 'date'),
        ('name', 'pubdate'),
        ('name', 'publish-date'),
        ('property', 'og:updated_time'),
    ]
    for attr, name in meta_tags:
        tag = soup.find('meta', attrs={attr: name})
        if tag and tag.get('content'):
            content = tag.get('content').strip()
            m = re.match(r'^(\d{4}-\d{2}-\d{2})', content)
            if m:
                try:
                    return datetime.strptime(m.group(1), "%Y-%m-%d")
                except ValueError:
                    pass
            parsed_dt = parse_natural_date(content)
            if parsed_dt:
                return parsed_dt

    # 4. Check HTML5 <time> tags
    for t in soup.find_all('time'):
        dt = t.get('datetime')
        if dt:
            m = re.match(r'^(\d{4}-\d{2}-\d{2})', dt.strip())
            if m:
                try:
                    return datetime.strptime(m.group(1), "%Y-%m-%d")
                except ValueError:
                    pass
        t_text = t.get_text().strip()
        parsed_dt = parse_natural_date(t_text)
        if parsed_dt:
            return parsed_dt

    # 5. Check URL path for date structure
    m = re.search(r'/(\d{4})/(\d{2})/(\d{2})/', url)
    if m:
        try:
            return datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            pass
            
    m = re.search(r'/(\d{4})/(\d{1,2})/(\d{1,2})/', url)
    if m:
        try:
            return datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            pass

    # 6. Check main content text
    main_content = get_main_content(soup)
    body_text = main_content.get_text()
    body_text_clean = re.sub(r'\s+', ' ', body_text)
    parsed_dt = parse_natural_date(body_text_clean)
    if parsed_dt:
        return parsed_dt

    return None

def check_date_matches(soup, url, target_date):
    """
    Checks if the web page indicates it was published or updated on target_date.
    """
    year = str(target_date.year)
    month_pad = target_date.strftime("%m")
    month_unpad = str(target_date.month)
    day_pad = target_date.strftime("%d")
    day_unpad = str(target_date.day)
    
    # 1. Check URL path
    url_patterns = [
        f"/{year}/{month_pad}/{day_pad}/",
        f"/{year}/{month_unpad}/{day_unpad}/",
        f"/{year}-{month_pad}-{day_pad}",
        f"/{year}/{month_pad}/{day_pad}",
    ]
    for pattern in url_patterns:
        if pattern in url:
            return True
            
    # 2. Check meta tags
    iso_date = target_date.strftime("%Y-%m-%d")
    meta_tags = [
        ('property', 'article:published_time'),
        ('property', 'article:modified_time'),
        ('name', 'date'),
        ('name', 'pubdate'),
        ('name', 'publish-date'),
        ('property', 'og:updated_time'),
    ]
    for attr, name in meta_tags:
        tag = soup.find('meta', attrs={attr: name})
        if tag and tag.get('content'):
            content = tag.get('content')
            if iso_date in content:
                return True
                
    # Check itemprop date metadata on any tag (e.g. meta, span, time)
    for tag in soup.find_all(attrs={'itemprop': re.compile(r'date', re.IGNORECASE)}):
        content = tag.get('content') or tag.get_text()
        if content and iso_date in content:
            return True
            
    # 3. Check HTML5 <time> tags
    for t in soup.find_all('time'):
        dt = t.get('datetime')
        if dt and iso_date in dt:
            return True
            
    # 4. Check JSON-LD
    for script in soup.find_all('script', type='application/ld+json'):
        if not script.string:
            continue
        try:
            data = json.loads(script.string)
            dates = []
            def find_dates(obj):
                if isinstance(obj, dict):
                    for k, v in obj.items():
                        if k in ['datePublished', 'dateModified', 'uploadDate']:
                            if isinstance(v, str):
                                dates.append(v)
                        else:
                            find_dates(v)
                elif isinstance(obj, list):
                    for item in obj:
                        find_dates(item)
            find_dates(data)
            for date_str in dates:
                if iso_date in date_str:
                    return True
        except Exception:
            pass

    # 5. Check page text, but restricted to the main content body to avoid sidebar false positives
    month_name = target_date.strftime("%B")
    month_abbr = target_date.strftime("%b")
    
    text_patterns = [
        rf"\b{month_name}\s+{day_unpad},?\s+{year}\b",
        rf"\b{month_name}\s+{day_pad},?\s+{year}\b",
        rf"\b{day_unpad}\s+{month_name}\s+{year}\b",
        rf"\b{day_pad}\s+{month_name}\s+{year}\b",
        rf"\b{month_abbr}\s+{day_unpad},?\s+{year}\b",
        rf"\b{month_abbr}\s+{day_pad},?\s+{year}\b",
        rf"\b{day_unpad}\s+{month_abbr}\s+{year}\b",
        rf"\b{day_pad}\s+{month_abbr}\s+{year}\b",
        rf"\b{iso_date}\b",
        rf"\b{target_date.strftime('%d/%m/%Y')}\b",
        rf"\b{target_date.strftime('%m/%d/%Y')}\b"
    ]
    
    main_content = get_main_content(soup)
    body_text = main_content.get_text()
    body_text_clean = re.sub(r'\s+', ' ', body_text)
    
    for pattern in text_patterns:
        if re.search(pattern, body_text_clean, re.IGNORECASE):
            return True
            
    return False

def check_published_today(soup, url):
    """
    Checks if the web page indicates it was published today.
    """
    today = datetime.now()
    
    # Check if matches today
    if check_date_matches(soup, url, today):
        return True
        
    # Check colloquial terms in <time> tags or class text (which are relative to today)
    for t in soup.find_all('time'):
        t_text = t.get_text().lower()
        if any(term in t_text for term in ["today", "hours ago", "hour ago", "mins ago", "minutes ago"]):
            return True
            
    for element in soup.find_all(['span', 'p', 'div', 'a']):
        classes = element.get('class', [])
        if any(any(ind in str(c).lower() for ind in ['date', 'time', 'meta', 'pub']) for c in classes):
            el_text = re.sub(r'\s+', ' ', element.get_text().lower())
            if len(el_text) < 150:
                if any(term in el_text for term in ["today", "hours ago", "hour ago", "mins ago"]):
                    return True
                    
    return False

def get_main_content(soup):
    """
    Attempts to extract the main article container from the soup, skipping sidebars.
    """
    selectors = [
        ('div', {'id': 'articlebody'}),
        ('div', {'id': 'article-body'}),
        ('div', {'class': 'post-body'}),
        ('div', {'class': 'article-content'}),
        ('div', {'class': 'entry-content'}),
        ('div', {'class': 'post'}),
        ('article', {'class': 'blog-post'}),
        ('article', {}),
        ('main', {}),
    ]
    for tag, attrs in selectors:
        element = soup.find(tag, attrs=attrs)
        if element:
            if len(element.get_text()) > 200:
                return element
    return soup

def determine_severity_and_teams(title, description, raw_text):
    """
    Heuristically determines vulnerability severity and affected teams from text content.
    """
    full_text = f"{title} {description} {raw_text}".lower()
    
    # 1. Determine Severity (Critical, High, Medium, Low)
    severity = "Medium" # Default fallback
    
    # Look for CVSS scores (e.g. CVSS 9.8, score of 8.5)
    cvss_matches = re.findall(r'\b(?:cvss|score|v3|v2)\s*(?:score)?\s*(?:of|:)?\s*([0-9]\.[0-9])\b', full_text)
    if cvss_matches:
        try:
            scores = [float(s) for s in cvss_matches if 0.0 <= float(s) <= 10.0]
            if scores:
                max_score = max(scores)
                if max_score >= 9.0:
                    severity = "Critical"
                elif max_score >= 7.0:
                    severity = "High"
                elif max_score >= 4.0:
                    severity = "Medium"
                else:
                    severity = "Low"
        except ValueError:
            pass
            
    # Check for explicit severity keywords if CVSS wasn't found or was fallback
    if severity == "Medium":
        if "critical" in full_text:
            severity = "Critical"
        elif "high" in full_text:
            severity = "High"
        elif "medium" in full_text or "moderate" in full_text:
            severity = "Medium"
        elif "low" in full_text or "minor" in full_text:
            severity = "Low"
            
    # 2. Determine Teams Affected
    TEAM_KEYWORDS = {
        'Windows': [r'\bwindows\b', r'\bmicrosoft\b', r'\bactive directory\b', r'\biis\b', r'\bexchange server\b', r'\bazure\b', r'\bwin32\b', r'\bwin64\b'],
        'Linux': [r'\blinux\b', r'\bubuntu\b', r'\bdebian\b', r'\bredhat\b', r'\brhel\b', r'\bcentos\b', r'\bunix\b', r'\bsuse\b', r'\bkernel\b', r'\bnginx\b', r'\bapache\b'],
        'Network': [r'\bnetwork\b', r'\bcisco\b', r'\bjuniper\b', r'\bfortinet\b', r'\bpalo alto\b', r'\brouter\b', r'\bswitch\b', r'\bfirewall\b', r'\bvpn\b', r'\bsd-wan\b', r'\bnetscaler\b', r'\bf5 big-ip\b', r'\bdns\b', r'\bdhcp\b'],
        'DISE': [r'\bdise\b', r'\bsap\b', r'\berp\b', r'\bcrm\b', r'\bsharepoint\b', r'\bsalesforce\b'],
        'Storage and Backup': [r'\bstorage\b', r'\bbackup\b', r'\bnas\b', r'\bsan\b', r'\bnetapp\b', r'\bveeam\b', r'\bcommvault\b', r'\brubrik\b', r'\bqnap\b', r'\bsynology\b', r'\bveritas\b'],
        'Command Center': [r'\bcommand center\b', r'\bnoc\b', r'\bsoc\b', r'\bmonitoring\b', r'\bsiem\b', r'\bsplunk\b', r'\bsolarwinds\b', r'\bdatadog\b', r'\bdynatrace\b', r'\bnagios\b'],
        'Desktop Support': [r'\bdesktop\b', r'\blaptop\b', r'\bchrome\b', r'\bfirefox\b', r'\bsafari\b', r'\bedge\b', r'\boffice\b', r'\bzoom\b', r'\badobe\b', r'\breader\b', r'\bmacos\b', r'\bwindows 10\b', r'\bwindows 11\b'],
        'Database': [r'\bdatabase\b', r'\bsql\b', r'\bmysql\b', r'\bpostgresql\b', r'\boracle\b', r'\bmssql\b', r'\bmongodb\b', r'\bredis\b', r'\bmariadb\b', r'\bsqlite\b', r'\bnosql\b']
    }
    
    matched_teams = []
    for team, patterns in TEAM_KEYWORDS.items():
        for pattern in patterns:
            if re.search(pattern, full_text):
                matched_teams.append(team)
                break # Move to next team
                
    if not matched_teams:
        matched_teams = ["Unknown"]
        
    teams_str = ", ".join(matched_teams)
    
    return severity, teams_str

def parse_article(url, html_content):
    """
    Parses HTML content, finds CVEs, and extracts metadata for each CVE found.
    Returns a list of dicts with date validation info.
    """
    soup = BeautifulSoup(html_content, 'html.parser')
    
    # 1. Extract Title
    title = ""
    og_title = soup.find('meta', attrs={'property': 'og:title'})
    if og_title and og_title.get('content'):
        title = og_title.get('content')
    
    if not title and soup.title:
        title = soup.title.string
        
    if not title:
        h1 = soup.find('h1')
        if h1:
            title = h1.get_text()
            
    title = clean_text(title) or "Untitled Article"

    # Get a general meta description for fallback
    meta_description = extract_meta_description(soup)

    # Find first descriptive paragraph for fallback if meta description is empty
    first_p = ""
    for p in soup.find_all('p'):
        text = clean_text(p.get_text())
        if len(text) > 60:
            first_p = text
            break

    # Determine if this article is published/modified today
    is_today = check_published_today(soup, url)

    # Extract article date
    published_date = extract_article_date(soup, url)
    if not published_date and is_today:
        published_date = datetime.now()

    # 2. Extract CVEs
    # Extract main content container to avoid sidebar/footer false positives
    main_content = get_main_content(soup)
    raw_text = main_content.get_text()
    found_cves = CVE_REGEX.findall(raw_text)
    
    # Normalize to uppercase and deduplicate
    unique_cves = sorted(list(set(cve.upper() for cve in found_cves)))
    
    results = []
    
    # If no CVEs are found, we return an empty list or we can log it
    if not unique_cves:
        return results

    # Determine severity and teams once for the article text
    severity, teams_str = determine_severity_and_teams(title, meta_description or first_p, raw_text)

    for cve in unique_cves:
        # Try to find paragraph mentioning this specific CVE
        description = find_cve_context(soup, cve)
        
        # Fallbacks for description
        if not description:
            description = meta_description
        if not description:
            description = first_p
        if not description:
            description = "Vulnerability details discussed in the article. No direct text context found."

        results.append({
            'CVE ID': cve,
            'Severity': severity,
            'Teams Affected': teams_str,
            'Title': title,
            'Source': url,
            'Description': description,
            'Published Today': is_today,
            'Date': published_date
        })
        
    return results

def generate_excel_report(data_list):
    """
    Creates an Excel spreadsheet byte stream from a list of dicts.
    Applies professional native table formatting to the sheet.
    """
    df = pd.DataFrame(data_list)
    
    # Standardize column structure and ordering
    cols = ['CVE ID', 'Severity', 'Teams Affected', 'Title', 'Source', 'Description']
    for c in cols:
        if c not in df.columns:
            df[c] = ""
    df = df[cols]
    
    # Reorder columns to have S.No first
    df.insert(0, 'S.No', range(1, len(df) + 1))

    # Create an in-memory buffer
    excel_buffer = io.BytesIO()
    
    with pd.ExcelWriter(excel_buffer, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Threat Intelligence')
        
        # Get worksheet for styling
        workbook = writer.book
        worksheet = writer.sheets['Threat Intelligence']
        
        # Define Excel Table object for proper table formatting with auto-filters
        last_col_letter = get_column_letter(len(df.columns))
        last_row_index = len(df) + 1
        
        tab = Table(displayName="ThreatIntelTable", ref=f"A1:{last_col_letter}{last_row_index}")
        
        # TableStyleMedium9 has clean corporate blue headers and automatic row zebra-striping
        style = TableStyleInfo(
            name="TableStyleMedium9", 
            showFirstColumn=False,
            showLastColumn=False, 
            showRowStripes=True, 
            showColumnStripes=False
        )
        tab.tableStyleInfo = style
        worksheet.add_table(tab)
        
        # Add custom fonts and alignment mapping
        data_font = Font(name='Segoe UI', size=10)
        cve_font = Font(name='Segoe UI', size=10, bold=True, color='C00000') # Accent red for CVE IDs
        
        align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
        align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
        
        worksheet.row_dimensions[1].height = 28
        
        # Style Data Rows (add alignment & font styling)
        for row_num in range(2, len(df) + 2):
            worksheet.row_dimensions[row_num].height = 24
            for col_num in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=row_num, column=col_num)
                cell.font = data_font
                
                # Alignments based on content types
                col_name = df.columns[col_num - 1]
                if col_name in ['S.No', 'CVE ID', 'Severity', 'Teams Affected']:
                    cell.alignment = align_center
                    if col_name == 'CVE ID':
                        cell.font = cve_font
                    elif col_name == 'Severity':
                        sev_val = str(cell.value)
                        if sev_val == "Critical":
                            cell.font = Font(name='Segoe UI', size=10, bold=True, color='C00000') # Dark Red
                        elif sev_val == "High":
                            cell.font = Font(name='Segoe UI', size=10, bold=True, color='E65100') # Dark Orange
                        elif sev_val == "Medium":
                            cell.font = Font(name='Segoe UI', size=10, bold=True, color='2563EB') # Blue
                        elif sev_val == "Low":
                            cell.font = Font(name='Segoe UI', size=10, bold=True, color='16A34A') # Green
                        else:
                            cell.font = Font(name='Segoe UI', size=10, italic=True, color='64748B')
                else:
                    cell.alignment = align_left
        
        # Set precise column widths for clean readability
        for col in worksheet.columns:
            col_name = col[0].value
            col_letter = get_column_letter(col[0].column)
            
            if col_name == 'S.No':
                worksheet.column_dimensions[col_letter].width = 8
            elif col_name == 'CVE ID':
                worksheet.column_dimensions[col_letter].width = 18
            elif col_name == 'Severity':
                worksheet.column_dimensions[col_letter].width = 12
            elif col_name == 'Teams Affected':
                worksheet.column_dimensions[col_letter].width = 22
            elif col_name == 'Title':
                worksheet.column_dimensions[col_letter].width = 35
            elif col_name == 'Source':
                worksheet.column_dimensions[col_letter].width = 40
            elif col_name == 'Description':
                worksheet.column_dimensions[col_letter].width = 55
            else:
                worksheet.column_dimensions[col_letter].width = 20

    excel_buffer.seek(0)
    return excel_buffer.getvalue()

def parse_cisa_kev(url):
    """
    Fetches the CISA KEV JSON feed and parses it.
    Returns findings based on the dateAdded field (last 30 days only to avoid bloating).
    """
    json_url = "https://www.cisa.gov/sites/default/files/feeds/known_exploited_vulnerabilities.json"
    headers = {
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
    }
    try:
        response = requests.get(json_url, headers=headers, verify=False, timeout=15)
        response.raise_for_status()
        data = response.json()
    except Exception as e:
        raise Exception(f"Failed to fetch CISA KEV JSON: {str(e)}")
        
    vulnerabilities = data.get('vulnerabilities', [])
    findings = []
    
    today = datetime.now()
    thirty_days_ago = today - timedelta(days=30)
    
    today_str = today.strftime("%Y-%m-%d")
    thirty_days_ago_str = thirty_days_ago.strftime("%Y-%m-%d")
    
    # Sort vulnerabilities by dateAdded descending
    sorted_vulns = sorted(vulnerabilities, key=lambda x: x.get('dateAdded', ''), reverse=True)
    
    for v in sorted_vulns:
        date_added = v.get('dateAdded', '')
        is_today = (date_added == today_str)
        
        # Keep only the last 30 days of KEV additions to avoid bloating reports
        if is_today or date_added >= thirty_days_ago_str:
            title = f"CISA KEV Catalog: {v.get('vulnerabilityName', '')}"
            desc = v.get('shortDescription', '')
            
            # Determine severity and teams heuristically
            severity, teams_str = determine_severity_and_teams(title, desc, "")
            
            # Convert date_added to datetime
            published_date = None
            if date_added:
                try:
                    published_date = datetime.strptime(date_added, "%Y-%m-%d")
                except ValueError:
                    pass
            
            findings.append({
                'CVE ID': v.get('cveID', ''),
                'Severity': severity,
                'Teams Affected': teams_str,
                'Title': title,
                'Source': url,
                'Description': desc,
                'Published Today': is_today,
                'Date': published_date
            })
            
    return findings

def group_findings(findings_list):
    """
    Groups findings by Title, Source, and Published Today.
    If multiple CVEs exist for the same article, merges their CVE IDs with a comma,
    takes the highest severity, combines affected teams, and joins descriptions.
    """
    grouped = {}
    for f in findings_list:
        key = (f['Title'], f['Source'], f['Published Today'])
        if key not in grouped:
            grouped[key] = {
                'CVEs': set(),
                'Descriptions': [],
                'Severities': set(),
                'Teams': set(),
                'Date': f.get('Date')
            }
        
        # Split and clean CVE IDs to handle both single and already-comma-separated ones
        cves = [c.strip() for c in f['CVE ID'].split(',')]
        for cve in cves:
            if cve:
                grouped[key]['CVEs'].add(cve)
                
        desc = f['Description']
        if desc and desc not in grouped[key]['Descriptions']:
            grouped[key]['Descriptions'].append(desc)
            
        # Add Severity and Teams Affected
        if f.get('Severity'):
            grouped[key]['Severities'].add(f['Severity'])
        if f.get('Teams Affected'):
            teams = [t.strip() for t in f['Teams Affected'].split(',')]
            for t in teams:
                if t:
                    grouped[key]['Teams'].add(t)
            
    result = []
    # Severity classification hierarchy
    sev_order = {"Critical": 4, "High": 3, "Medium": 2, "Low": 1, "Unknown": 0}
    
    for (title, source, published_today), data in grouped.items():
        sorted_cves = sorted(list(data['CVEs']))
        cve_str = ", ".join(sorted_cves)
        
        # Join descriptions with double newline for clean reading in wrapped cells
        combined_desc = "\n\n".join(data['Descriptions'])
        
        # Determine highest severity level present in grouped findings
        highest_sev = "Unknown"
        max_val = -1
        for s in data['Severities']:
            val = sev_order.get(s, 0)
            if val > max_val:
                max_val = val
                highest_sev = s
        if highest_sev == "Unknown" and not data['Severities']:
            highest_sev = "Medium" # Fallback
            
        # Deduplicate and sort affected teams (omit "Unknown" if other teams are identified)
        teams_list = list(data['Teams'])
        if len(teams_list) > 1 and "Unknown" in teams_list:
            teams_list.remove("Unknown")
        sorted_teams = sorted(teams_list)
        teams_str = ", ".join(sorted_teams) if sorted_teams else "Unknown"
        
        result.append({
            'CVE ID': cve_str,
            'Severity': highest_sev,
            'Teams Affected': teams_str,
            'Title': title,
            'Source': source,
            'Description': combined_desc,
            'Published Today': published_today,
            'Date': data['Date']
        })
    return result
